import openpyxl
from office365.sharepoint.client_context import ClientContext
from office365.sharepoint.listitems.list_item import ListItem
from office365.runtime.auth.client_credential import ClientCredential

# Excel file path
EXCEL_FILE = r"path\to\your\excel\file.xlsx"

# SharePoint credentials
SHAREPOINT_SITE_URL = "https://yourtenant.sharepoint.com/sites/mysharepoint"
CLIENT_ID = "your_client_id"
CLIENT_SECRET = "your_client_secret"

def read_excel_table(file_path, table_name):
    """Read data from an Excel table."""
    workbook = openpyxl.load_workbook(file_path)
    worksheet = workbook.active
    
    data = []
    table = None
    
    # Find the table by name
    for tbl in worksheet.tables.values():
        if tbl.name == table_name:
            table = tbl
            break
    
    if not table:
        print(f"Table '{table_name}' not found in the workbook.")
        return data
    
    # Extract table boundaries
    min_row = table.ref.split(':')[0]
    min_row = int(''.join(filter(str.isdigit, min_row)))
    max_row = table.ref.split(':')[1]
    max_row = int(''.join(filter(str.isdigit, max_row)))
    
    # Read rows from the table
    for row in worksheet.iter_rows(min_row=min_row, max_row=max_row, values_only=True):
        data.append(row)
    
    return data

def populate_sharepoint_list(data, site_url, client_id, client_secret, list_name):
    """Populate SharePoint list with data from Excel."""
    try:
        # Authenticate to SharePoint
        credentials = ClientCredential(client_id, client_secret)
        ctx = ClientContext(site_url).with_credentials(credentials)
        
        # Get the list
        sp_list = ctx.web.lists.get_by_title(list_name)
        
        # Skip header row if present
        for index, row in enumerate(data[1:], start=1):
            try:
                # Create list item properties
                item_properties = {}
                if row:
                    # Map Excel columns to SharePoint list columns
                    # Adjust column names based on your actual SharePoint list columns
                    item_properties['Title'] = row[0] if row[0] else f"Item {index}"
                    
                    # Add more field mappings as needed
                    for col_index, value in enumerate(row[1:], start=1):
                        item_properties[f'Column{col_index}'] = value
                
                # Add item to SharePoint list
                item = sp_list.add_item(item_properties)
                ctx.execute_query()
                print(f"Row {index} added successfully.")
                
            except Exception as e:
                print(f"Error adding row {index}: {str(e)}")
        
        print(f"All {len(data) - 1} rows have been processed.")
        
    except Exception as e:
        print(f"Error connecting to SharePoint: {str(e)}")

def main():
    """Main function to orchestrate the process."""
    print("Reading Excel table...")
    table_data = read_excel_table(EXCEL_FILE, "ReXTable")
    
    if table_data:
        print(f"Found {len(table_data)} rows in the table.")
        print("Populating SharePoint list...")
        populate_sharepoint_list(table_data, SHAREPOINT_SITE_URL, CLIENT_ID, CLIENT_SECRET, "mysharepointlist")
    else:
        print("No data found in Excel table.")

if __name__ == "__main__":
    main()
